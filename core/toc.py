"""Word 目录：静态扫描标题（纯 python-docx，不启动 Word）。

这里只做"读"，不写任何文件。它的作用是让界面在动手之前就知道
"这份文档有几个标题、分别是几级"，没有标题时直接拦住 ——
否则 Word 会给你插一个只有「未找到目录项。」的空目录（实测不报错）。

真正插入目录在 core.office.add_toc_to_docx（必须借 Word 算页码）。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOCX_EXTS = {".docx"}          # v1 只接 .docx；.doc/.rtf 请用户先另存
MAX_LEVEL = 4

# styles.xml 里 w:name 是规范化英文名，中文 Word 只影响 COM 的 NameLocal，
# 所以 "Heading 1" 才是主路径；「标题 1」留作兜底（手工改过样式的文档）。
_NAME_RE = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.IGNORECASE)
_ID_RE = re.compile(r"^Heading([1-9])$", re.IGNORECASE)


def _ppr_level(ppr) -> int | None:
    """从 <w:pPr> 里读 <w:outlineLvl>；0..8 对应 1..9 级，9 表示正文（返回 None）。"""
    if ppr is None:
        return None
    try:
        node = ppr.find(qn("w:outlineLvl"))
    except Exception:
        return None
    if node is None:
        return None
    val = node.get(qn("w:val"))
    if not (val or "").isdigit():
        return None
    n = int(val)
    return (n + 1) if n < 9 else None


def _style_level(style, depth: int = 0) -> int | None:
    """沿样式继承链找级别：outlineLvl → styleId → 名字 → base_style 递归。"""
    if style is None or depth > 8:
        return None
    try:
        el = style.element
    except Exception:
        el = None
    if el is not None:
        level = _ppr_level(el.find(qn("w:pPr")))
        if level is not None:
            return level

    m = _ID_RE.match(str(getattr(style, "style_id", "") or ""))
    if m:
        return int(m.group(1))
    m = _NAME_RE.match(str(getattr(style, "name", "") or "").strip())
    if m:
        return int(m.group(1))

    try:
        base = style.base_style
    except Exception:
        base = None
    return _style_level(base, depth + 1)


def level_of(p) -> int | None:
    """段落的大纲级别；None 表示正文（或无法判定）。"""
    try:
        ppr = p._p.pPr
    except Exception:
        ppr = None
    level = _ppr_level(ppr)
    if level is not None:
        return level
    try:
        style = p.style
    except Exception:
        style = None
    return _style_level(style)


def scan_headings(path: str | Path, max_level: int = MAX_LEVEL) -> list[tuple[int, str]]:
    """返回 [(级别, 文本)]，按文档顺序。文档读不了时抛 ValueError。

    已知边界：doc.paragraphs 不含表格单元格里的段落，所以表格内的标题扫不到
    （Word 自己插目录时会收录）。这里宁少不多，避免界面统计与 Word 结果打架。
    """
    path = Path(path)
    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"读不了这份文档：{exc}") from exc

    rows: list[tuple[int, str]] = []
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        level = level_of(p)
        if level is None or level > max_level:
            continue
        rows.append((level, text))
    return rows


def level_counts(rows: list[tuple[int, str]], depth: int = 3) -> dict[int, int]:
    """{1: 一级标题数, 2: …, …}，只统计 <= depth 的。"""
    depth = max(1, int(depth))
    out = {i: 0 for i in range(1, depth + 1)}
    for level, _text in rows:
        if 1 <= level <= depth:
            out[level] += 1
    return out


def preview_lines(rows: list[tuple[int, str]], depth: int = 3, limit: int = 200) -> list[str]:
    """缩进大纲，直接给界面日志框 join('\\n')。"""
    lines: list[str] = []
    for level, text in rows:
        if level > depth:
            continue
        indent = "    " * (level - 1)
        mark = "└ " if level > 1 else ""
        lines.append(f"{indent}{mark}{text}")
        if len(lines) >= limit:
            lines.append(f"…… 后面还有，只显示前 {limit} 条")
            break
    return lines


# 目录域的特征：instrText 通常是 ' TOC \o "1-3" \h \z \u ' 这种形状。
# 只认开头就是 TOC 的，避免正文里恰好提到这个词造成误判。
_TOC_INSTR_RE = re.compile(r"^\s*TOC\b", re.IGNORECASE)


def has_toc(path) -> bool:
    """文档里是否已经有目录。

    两种都算：Word 原生插入的 TOC 域，以及 WPS 那种把整个目录包在
    <w:sdt> 内容控件里的写法 —— 后者在 OOXML 层面里面装的还是同一套 TOC 域，
    所以遍历 instrText 就能一起认出来。

    这个判断只影响界面上的提示文案和 core.office 的处理路径：
    有目录 → 就地更新，位置不动；没目录 → 才新建。
    """
    try:
        doc = Document(str(path))
    except Exception:
        return False
    body = doc.element.body
    try:
        for node in body.iter(qn("w:instrText")):
            if _TOC_INSTR_RE.search(node.text or ""):
                return True
        for node in body.iter(qn("w:fldSimple")):
            if _TOC_INSTR_RE.search(node.get(qn("w:instr")) or ""):
                return True
    except Exception:
        return False
    return False
