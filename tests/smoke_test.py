"""无界面自测：解析样例文件，走完所有导出格式，并回读产物做校验。

运行： python tests/smoke_test.py
"""
from __future__ import annotations

import hashlib
import shutil
import sys
import time
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pymupdf  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

from core import office, outline, toc  # noqa: E402
from core.exporter import run_jobs, run_merge  # noqa: E402
from core.models import ExportSpec, Job, PageRef  # noqa: E402
from core.sources import open_source  # noqa: E402

OUT = ROOT / "out_smoke"
SAMPLES = ROOT / "samples"

_results: list[tuple[bool, str, str]] = []


def check(label: str, fn) -> None:
    t0 = time.time()
    try:
        detail = fn() or ""
        _results.append((True, label, f"{detail} ({time.time() - t0:.1f}s)"))
    except Exception as exc:  # noqa: BLE001
        _results.append((False, label, f"{exc}"))
        traceback.print_exc()


def spec(fmt: str, **kw) -> ExportSpec:
    base = dict(out_dir=OUT / fmt, dpi=120, overwrite="overwrite")
    base.update(kw)
    return ExportSpec(fmt=fmt, **base)


def pick(src, keys: list[str]):
    units = [u for u in src.units() if u.key in set(keys)]
    assert len(units) == len(keys), f"缺少单位：{keys}"
    return Job(source=src, units=units)


def main() -> int:
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True, exist_ok=True)

    if not (SAMPLES / "样例报告.pdf").exists():
        print("缺少样例文件，先运行： python tools/make_samples.py")
        return 2

    # ---------------------------------------------------------------- PDF
    pdf_src = open_source(SAMPLES / "样例报告.pdf")
    print(f"PDF 单位数 = {len(pdf_src.units())}")

    def pdf_thumb() -> str:
        img = pdf_src.thumbnail("p1", 320)
        img.save(OUT / "preview_pdf.png")
        assert img.width > 100
        return f"{img.width}x{img.height}"

    def pdf_png() -> str:
        files = run_jobs([pick(pdf_src, ["p1", "p3", "p5"])], spec("png"))
        assert len(files) == 3, files
        for f in files:
            with Image.open(f) as im:
                assert im.width > 800 and im.height > 1000, (f, im.size)
        return f"{len(files)} 张图 {im.size[0]}x{im.size[1]}"

    def pdf_jpg_long() -> str:
        files = run_jobs([pick(pdf_src, ["p1", "p2", "p3"])], spec("jpg", long_image=True))
        assert len(files) == 1, files
        with Image.open(files[0]) as im:
            assert im.height > im.width * 2
        return f"长图 {im.width}x{im.height}"

    def pdf_pdf() -> str:
        files = run_jobs([pick(pdf_src, ["p1", "p2", "p5"])], spec("pdf"))
        doc = pymupdf.open(files[0])
        try:
            assert doc.page_count == 3, doc.page_count
            return f"{doc.page_count} 页"
        finally:
            doc.close()

    def pdf_docx() -> str:
        files = run_jobs([pick(pdf_src, ["p2", "p4"])], spec("docx", dpi=100))
        from docx import Document

        d = Document(str(files[0]))
        assert len(d.inline_shapes) == 2, len(d.inline_shapes)
        return f"{len(d.inline_shapes)} 张图已排入"

    # ---------------------------------------------------------------- Word
    word_src = open_source(SAMPLES / "样例方案.docx")
    print(f"Word 单位数 = {len(word_src.units())}")

    def word_png() -> str:
        files = run_jobs([pick(word_src, ["p1", "p2"])], spec("png"))
        with Image.open(files[0]) as im:
            return f"{len(files)} 张图 {im.size[0]}x{im.size[1]}"

    def word_pdf() -> str:
        files = run_jobs([pick(word_src, ["p1", "p3"])], spec("pdf"))
        doc = pymupdf.open(files[0])
        try:
            assert doc.page_count == 2
            text = doc[0].get_text()
            assert "实施方案" in text or len(text) > 20, text[:60]
            return f"{doc.page_count} 页，首页含文字 {len(text)} 字符"
        finally:
            doc.close()

    def word_docx() -> str:
        files = run_jobs([pick(word_src, ["p1"])], spec("docx", dpi=110))
        return f"{files[0].name} {files[0].stat().st_size // 1024} KB"

    # ---------------------------------------------------------------- Excel
    xls_src = open_source(SAMPLES / "样例数据.xlsx")
    print(f"Excel 单位数 = {len(xls_src.units())} -> {[u.label for u in xls_src.units()]}")

    def xls_thumb() -> str:
        img = xls_src.thumbnail("s::销售明细", 320)
        img.save(OUT / "preview_xlsx.png")
        return f"{img.width}x{img.height}"

    def xls_png() -> str:
        files = run_jobs([pick(xls_src, ["s::销售明细"])], spec("png", dpi=150))
        with Image.open(files[0]) as im:
            assert im.width > 800, im.size
            return f"{len(files)} 张图 {im.size[0]}x{im.size[1]}"

    def xls_xlsx() -> str:
        files = run_jobs([pick(xls_src, ["s::销售明细", "s::区域汇总"])], spec("xlsx"))
        wb = load_workbook(files[0])
        try:
            assert wb.sheetnames == ["销售明细", "区域汇总"], wb.sheetnames
            assert wb["销售明细"]["A2"].value == "订单号"
            return f"工作表 {wb.sheetnames}"
        finally:
            wb.close()

    def xls_xlsx_separate() -> str:
        files = run_jobs([pick(xls_src, ["s::区域汇总", "s::使用说明"])], spec("xlsx", merge_xlsx=False))
        assert len(files) == 2, files
        return f"{len(files)} 个工作簿"

    def xls_pdf() -> str:
        files = run_jobs([pick(xls_src, ["s::区域汇总"])], spec("pdf"))
        doc = pymupdf.open(files[0])
        try:
            assert doc.page_count >= 1
            return f"{doc.page_count} 页"
        finally:
            doc.close()

    # ---------------------------------------------------------------- PPT
    ppt_src = open_source(SAMPLES / "样例演示.pptx")
    print(f"PPT 单位数 = {len(ppt_src.units())}")

    def ppt_thumb() -> str:
        img = ppt_src.thumbnail("p1", 320)
        img.save(OUT / "preview_pptx.png")
        assert img.width > 100, img.size
        return f"{img.width}x{img.height}"

    def ppt_png() -> str:
        files = run_jobs([pick(ppt_src, ["p1", "p3"])], spec("png"))
        assert len(files) == 2, files
        for f in files:
            with Image.open(f) as im:
                assert im.width > 800, (f.name, im.size)
        return f"{len(files)} 张图 {im.size[0]}x{im.size[1]}"

    def ppt_pdf() -> str:
        files = run_jobs([pick(ppt_src, ["p2", "p3", "p4"])], spec("pdf"))
        doc = pymupdf.open(files[0])
        try:
            assert doc.page_count == 3, doc.page_count
            text = doc[0].get_text()
            assert "季度经营摘要" in text, text[:80]
            return f"{doc.page_count} 页，首页文字 {len(text)} 字符"
        finally:
            doc.close()

    def ppt_pptx() -> str:
        import zipfile

        files = run_jobs([pick(ppt_src, ["p1", "p4"])], spec("pptx"))
        with zipfile.ZipFile(files[0]) as z:
            n = len([
                x for x in z.namelist()
                if x.startswith("ppt/slides/slide") and x.endswith(".xml")
            ])
        assert n == 2, f"子集里应有 2 页，实际 {n}"
        return f"{n} 页幻灯片 {files[0].stat().st_size // 1024} KB"

    # ---------------------------------------------------------------- Word 目录
    toc_out = OUT / "toc"
    contract = SAMPLES / "样例合同.docx"

    def toc_scan() -> str:
        rows = toc.scan_headings(SAMPLES / "样例方案.docx")
        counts = toc.level_counts(rows)
        assert len(rows) == 6, rows
        assert counts == {1: 3, 2: 3, 3: 0}, counts
        rows2 = toc.scan_headings(contract)
        assert len(rows2) == 8, rows2
        lines = toc.preview_lines(rows2, depth=2)
        assert any("└" in ln for ln in lines), lines       # 二级要缩进
        return f"方案 6 条 {counts}；合同 8 条"

    def toc_insert() -> str:
        before = hashlib.md5(contract.read_bytes()).hexdigest()
        target = toc_out / "样例合同_带目录.docx"
        res = office.add_toc_to_docx(contract, target, levels=3)
        assert res["entries"] == 8, res
        assert res["mode"] == "inserted", res        # 原文档没目录 → 走新建
        assert hashlib.md5(contract.read_bytes()).hexdigest() == before, "源文件被改动了！"

        d = Document(str(res["out"]))
        assert d.paragraphs[0].text.strip() == "目录", d.paragraphs[0].text
        assert "第 1 章 合同条款" in d.paragraphs[1].text, d.paragraphs[1].text[:60]

        # 幂等：对已带目录的文档再跑一次，应**就地更新**旧目录而不是删了重建
        # （重建要重新算插入点，一旦算错就会把目录挪到封面之前 —— 真实踩过的坑）
        again = office.add_toc_to_docx(res["out"], toc_out / "样例合同_再生成.docx", levels=3)
        assert again["mode"] == "updated", again
        assert again["replaced"] == 1, again
        assert again["entries"] == res["entries"], again     # 条目没变 = 没叠一份新的

        d2 = Document(str(again["out"]))
        assert d2.paragraphs[0].text.strip() == "目录", d2.paragraphs[0].text
        assert "第 1 章 合同条款" in d2.paragraphs[1].text, d2.paragraphs[1].text[:60]
        heads = sum(1 for p in d2.paragraphs if p.text.strip() == "目录")
        assert heads == 1, f"出现了 {heads} 个「目录」行，应该只有 1 个"

        p1 = pymupdf.open(office.word_to_pdf(target, force=True))
        p2 = pymupdf.open(office.word_to_pdf(again["out"], force=True))
        c1, c2 = p1.page_count, p2.page_count
        p1.close()
        p2.close()
        assert c1 == 5, c1
        assert c2 == 5, f"重复生成页数变了：{c2}"
        return f"{res['entries']} 条目录，{c1} 页，重复生成就地更新不叠加"

    def toc_merge() -> str:
        plan = open_source(SAMPLES / "样例方案.docx")
        cont = open_source(contract)
        try:
            pages = [
                PageRef(plan, "p1", "第 1 页", "样例方案.docx"),
                PageRef(cont, "p2", "第 2 页", "样例合同.docx"),
                PageRef(plan, "p3", "第 3 页", "样例方案.docx"),
            ]
            files = run_merge(pages, spec("docx", toc_levels=3))
            pdf = pymupdf.open(office.word_to_pdf(files[0].resolve(), force=True))
            try:
                pages_n = pdf.page_count
                first = pdf[0].get_text()
                second = pdf[1].get_text().strip()
                assert pages_n == 4, pages_n                      # 3 页正文 + 1 页目录
                assert "第 2 章 合同条款" in first, first[:150]
                assert "第 1 章 项目实施方案" in first, first[:150]
                assert second.startswith("第 1 章 项目实施方案"), second[:40]   # 正文页序没变
                return f"{pages_n} 页（首页目录，正文页序不变）"
            finally:
                pdf.close()
        finally:
            plan.close()
            cont.close()

    # ---------------------------------------------------------------- 手动标标题
    def _plain_docx(dst: Path) -> Path:
        """造一份「标题全是手打的」文档：没有标题样式，只有字号和加粗。

        真实场景就是这样 —— 用 WPS 写的人把章节标题写成正文 + 加粗 + 大字，
        Word 的目录一条都认不出来。样例目录里没有这种文件，只能现场造。
        """
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        d = Document()
        cover = d.add_paragraph()
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cover.add_run("校园闲置物品循环利用服务平台创业计划书")
        run.bold = True
        run.font.size = Pt(22)

        plan = [
            (1, "1 项目简介"),
            (2, "1.1 项目介绍"),
            (3, "1.1.1 平台定位"),
            (0, "本平台定位于校园专属的闲置物品循环利用服务平台，聚焦高校学生群体。" * 3),
            (3, "1.1.2 服务内容"),
            (0, "平台核心服务包含三大板块：闲置交易、租赁和捐赠，覆盖全场景。" * 3),
            (2, "1.2 项目背景"),
            (1, "2 市场分析"),
            (2, "2.1 国内"),
            (0, "近年来国内校园闲置经济市场规模呈逐年增长趋势，2024 年数据。" * 3),
        ]
        for level, text in plan:
            para = d.add_paragraph()
            run = para.add_run(text)
            if level:
                run.bold = True
                run.font.size = Pt({1: 16, 2: 14, 3: 12}[level])
        dst.parent.mkdir(parents=True, exist_ok=True)
        d.save(str(dst))
        return dst

    def outline_guess() -> str:
        plain = _plain_docx(toc_out / "手打标题.docx")
        assert len(toc.scan_headings(plain)) == 0, "这份文档不该有标题样式"

        paras = outline.list_paragraphs(plain)
        assert len(paras) == 11, len(paras)
        marks = outline.guess_levels(paras)

        # 编号规律是最可靠的判据：1 / 1.1 / 1.1.1 分别是 1/2/3 级
        expect = {1: 1, 2: 2, 3: 3, 5: 3, 7: 2, 8: 1, 9: 2}
        got = {i: marks.get(i, 0) for i in expect}
        assert got == expect, (got, expect)
        # 长正文不该被猜成标题
        for i in (4, 6, 10):
            assert marks.get(i, 0) == 0, (i, marks.get(i))
        return f"{len(paras)} 段，猜出 {len(marks)} 条，级别全对"

    def outline_apply() -> str:
        plain = _plain_docx(toc_out / "手打标题.docx")
        paras = outline.list_paragraphs(plain)
        marks = outline.guess_levels(paras)

        before = hashlib.md5(plain.read_bytes()).hexdigest()
        dst = toc_out / "手打标题_已标记.docx"
        n = outline.apply_levels(plain, dst, marks)
        assert n == len(marks), (n, len(marks))
        assert hashlib.md5(plain.read_bytes()).hexdigest() == before, "源文件被改动了！"

        from docx.oxml.ns import qn

        d = Document(str(dst))
        levels = {}
        for i, p in enumerate(d.paragraphs):
            ppr = p._p.find(qn("w:pPr"))
            node = None if ppr is None else ppr.find(qn("w:outlineLvl"))
            if node is not None:
                levels[i] = int(node.get(qn("w:val"))) + 1
        assert levels.get(1) == 1 and levels.get(3) == 3, levels

        # 最要紧的一条：外观一个字都不许变
        src_doc = Document(str(plain))
        for i, (a, b) in enumerate(zip(src_doc.paragraphs, d.paragraphs)):
            assert a.style.name == b.style.name, f"第 {i} 段样式被改了"
            assert a.alignment == b.alignment, f"第 {i} 段对齐被改了"
            for ra, rb in zip(a.runs, b.runs):
                assert ra.font.size == rb.font.size, f"第 {i} 段字号被改了"
                assert ra.font.bold == rb.font.bold, f"第 {i} 段加粗被改了"
        return f"{n} 段写了大纲级别，字体/字号/对齐均未变"

    def toc_marks() -> str:
        plain = _plain_docx(toc_out / "手打标题.docx")
        assert len(toc.scan_headings(plain)) == 0
        paras = outline.list_paragraphs(plain)
        marks = outline.guess_levels(paras)
        marks.pop(0, None)          # 封面大标题不该进目录

        out = toc_out / "手打标题_带目录.docx"
        res = office.add_toc_to_docx(plain, out, levels=3, marks=marks)
        assert res["mode"] == "inserted", res
        assert res["entries"] == len(marks), (res["entries"], len(marks))

        # 顺序必须是：封面 → 目录 → 正文。目录跑到封面前面是 v1.5.0 的著名 bug
        d = Document(str(out))
        assert d.paragraphs[0].text.strip().startswith("校园闲置"), d.paragraphs[0].text[:40]
        assert d.paragraphs[1].text.strip() == "目录", d.paragraphs[1].text
        assert "1 项目简介" in d.paragraphs[2].text, d.paragraphs[2].text[:40]

        # 有目录之后再生成：走就地更新，条目数不变（不能叠一份新的）
        # 注意 marks 要按**新文件**的段落下标重算 —— 插入目录后段落位置全变了，
        # 拿旧下标去标会标到目录那一行上（第一次写这个测试就栽在这儿）
        paras2 = outline.list_paragraphs(out)
        marks2 = outline.guess_levels(paras2)
        marks2.pop(0, None)                      # 封面大标题同样不进目录
        again = office.add_toc_to_docx(out, toc_out / "手打标题_再生成.docx",
                                       levels=3, marks=marks2)
        assert again["mode"] == "updated", again
        assert again["entries"] == res["entries"], (again["entries"], res["entries"])

        d3 = Document(str(again["out"]))
        heads = sum(1 for p in d3.paragraphs if p.text.strip() == "目录")
        assert heads == 1, f"出现了 {heads} 个「目录」行，应该只有 1 个"
        return f"{res['entries']} 条目录（纯大纲级别），重复生成不叠加"


    check("PDF  缩略图", pdf_thumb)
    check("PDF  → PNG（非连续页 p1/p3/p5）", pdf_png)
    check("PDF  → JPG（三页拼长图）", pdf_jpg_long)
    check("PDF  → PDF（重排为 3 页）", pdf_pdf)
    check("PDF  → DOCX（图片按页排版）", pdf_docx)
    check("Word 缩略图 / 解析分页", lambda: f"{len(word_src.units())} 页")
    check("Word → PNG", word_png)
    check("Word → PDF（保留矢量文字）", word_pdf)
    check("Word → DOCX", word_docx)
    check("Excel 缩略图（工作表渲染）", xls_thumb)
    check("Excel → PNG（整表一图）", xls_png)
    check("Excel → XLSX（多表合并）", xls_xlsx)
    check("Excel → XLSX（一表一文件）", xls_xlsx_separate)
    check("Excel → PDF", xls_pdf)
    check("PPT  缩略图（幻灯片渲染）", ppt_thumb)
    check("PPT  → PNG", ppt_png)
    check("PPT  → PDF（保留矢量文字）", ppt_pdf)
    check("PPT  → PPTX（所选幻灯片子集）", ppt_pptx)
    check("Word 目录：静态扫描标题样式", toc_scan)
    check("Word 目录：插入目录且不动源文件", toc_insert)
    check("Word 目录：合并文档自动插目录", toc_merge)
    check("手动标标题：自动预判级别", outline_guess)
    check("手动标标题：只写大纲级别，外观不变", outline_apply)
    check("手动标标题：标完能生成目录", toc_marks)

    print("\n" + "=" * 64)
    ok = sum(1 for r in _results if r[0])
    for passed, label, detail in _results:
        print(f"{'✔ PASS' if passed else '✘ FAIL'}  {label:34s} {detail}")
    print("=" * 64)
    print(f"合计 {ok}/{len(_results)} 通过；产物目录：{OUT}")
    total = sum(f.stat().st_size for f in OUT.rglob("*") if f.is_file())
    print(f"产物 {sum(1 for f in OUT.rglob('*') if f.is_file())} 个文件，{total / 1024:.0f} KB")
    return 0 if ok == len(_results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
